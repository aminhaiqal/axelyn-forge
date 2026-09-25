"""Single-column, ATS-friendly resume templates rendered from structured content."""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


TEMPLATE_VERSION = "1"
DEFAULT_TEMPLATE_ID = "ats-classic"


@dataclass(frozen=True)
class TemplateSpec:
    id: str
    name: str
    description: str
    density: str
    font: str
    body_size: float
    name_size: float
    heading_size: float
    accent: str
    margin: float
    paragraph_after: float


TEMPLATES = (
    TemplateSpec(
        id="ats-classic",
        name="Classic ATS",
        description="Traditional hierarchy with generous spacing and a centered identity block.",
        density="comfortable",
        font="Arial",
        body_size=10.5,
        name_size=24,
        heading_size=11.5,
        accent="172A35",
        margin=0.72,
        paragraph_after=3.5,
    ),
    TemplateSpec(
        id="ats-modern",
        name="Modern ATS",
        description="Crisp left-aligned structure with restrained color and clear section rules.",
        density="balanced",
        font="Aptos",
        body_size=10.25,
        name_size=25,
        heading_size=11.5,
        accent="315E73",
        margin=0.68,
        paragraph_after=3,
    ),
    TemplateSpec(
        id="ats-compact",
        name="Compact ATS",
        description="Tighter spacing for experienced candidates who need more room without columns.",
        density="compact",
        font="Arial",
        body_size=9.5,
        name_size=22,
        heading_size=10.5,
        accent="223A48",
        margin=0.55,
        paragraph_after=2,
    ),
)
TEMPLATE_BY_ID = {template.id: template for template in TEMPLATES}


def template_catalog() -> list[dict[str, str]]:
    return [
        {
            "id": template.id,
            "name": template.name,
            "description": template.description,
            "density": template.density,
        }
        for template in TEMPLATES
    ]


def _color(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def _font(run, spec: TemplateSpec, *, size: float | None = None, bold: bool = False) -> None:
    run.font.name = spec.font
    run.font.size = Pt(size or spec.body_size)
    run.font.bold = bold
    run.font.color.rgb = _color("172A35")
    run._element.rPr.rFonts.set(qn("w:eastAsia"), spec.font)


def _bottom_border(paragraph, color: str, size: str = "8") -> None:
    properties = paragraph._p.get_or_add_pPr()
    borders = properties.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        properties.append(borders)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), color)
    borders.append(bottom)


def _clean_lines(value: object) -> list[str]:
    if not isinstance(value, list):
        return []
    return [str(line).strip() for line in value if str(line).strip()]


def _add_section_heading(document: Document, title: str, spec: TemplateSpec) -> None:
    paragraph = document.add_paragraph()
    paragraph.style = document.styles["Resume Section"]
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run(title.upper())
    _font(run, spec, size=spec.heading_size, bold=True)
    run.font.color.rgb = _color(spec.accent)
    _bottom_border(paragraph, spec.accent, "6")


def _add_lines(document: Document, lines: list[str], spec: TemplateSpec) -> None:
    for line in lines:
        is_bullet = line.startswith(("•", "-", "–", "*"))
        value = line.lstrip("•-–* ").strip() if is_bullet else line
        if not value:
            continue
        paragraph = document.add_paragraph(style="Resume Bullet" if is_bullet else "Resume Body")
        run = paragraph.add_run(f"• {value}" if is_bullet else value)
        _font(run, spec, bold=not is_bullet and (" | " in value or " — " in value))


def render_resume(
    *,
    draft: dict[str, object],
    output: Path,
    title: str,
    description: str,
) -> tuple[str, str]:
    """Render every modeled section into a one-column DOCX without ATS-hostile objects."""
    template_id = str(draft.get("template_id") or DEFAULT_TEMPLATE_ID)
    spec = TEMPLATE_BY_ID.get(template_id, TEMPLATE_BY_ID[DEFAULT_TEMPLATE_ID])
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(spec.margin)
    section.bottom_margin = Inches(spec.margin)
    section.left_margin = Inches(spec.margin)
    section.right_margin = Inches(spec.margin)

    normal = document.styles["Normal"]
    normal.font.name = spec.font
    normal.font.size = Pt(spec.body_size)
    normal.font.color.rgb = _color("172A35")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), spec.font)

    body_style = document.styles.add_style("Resume Body", WD_STYLE_TYPE.PARAGRAPH)
    body_style.base_style = normal
    body_style.paragraph_format.space_after = Pt(spec.paragraph_after)
    body_style.paragraph_format.line_spacing = 1.04

    bullet_style = document.styles.add_style("Resume Bullet", WD_STYLE_TYPE.PARAGRAPH)
    bullet_style.base_style = normal
    bullet_style.paragraph_format.left_indent = Inches(0.18)
    bullet_style.paragraph_format.first_line_indent = Inches(-0.13)
    bullet_style.paragraph_format.space_after = Pt(spec.paragraph_after)
    bullet_style.paragraph_format.line_spacing = 1.02

    section_style = document.styles.add_style("Resume Section", WD_STYLE_TYPE.PARAGRAPH)
    section_style.base_style = normal
    section_style.paragraph_format.space_before = Pt(8 if spec.density != "compact" else 5)
    section_style.paragraph_format.space_after = Pt(4)

    identity = document.add_paragraph()
    identity.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
        if spec.id == "ats-classic"
        else WD_ALIGN_PARAGRAPH.LEFT
    )
    identity.paragraph_format.space_after = Pt(1)
    run = identity.add_run(str(draft.get("full_name") or "Your name"))
    _font(run, spec, size=spec.name_size, bold=True)
    run.font.color.rgb = _color(spec.accent)

    headline_value = str(draft.get("headline") or draft.get("target_role") or "").strip()
    if headline_value:
        headline = document.add_paragraph()
        headline.alignment = identity.alignment
        headline.paragraph_format.space_after = Pt(2)
        run = headline.add_run(headline_value)
        _font(run, spec, size=spec.body_size + 1, bold=True)

    contact_value = str(draft.get("contact_line") or "").strip()
    if contact_value:
        contact = document.add_paragraph()
        contact.alignment = identity.alignment
        contact.paragraph_format.space_after = Pt(7)
        run = contact.add_run(contact_value)
        _font(run, spec, size=max(spec.body_size - 0.5, 8.5))
        _bottom_border(contact, spec.accent, "5")

    summary = str(draft.get("summary") or "").strip()
    if summary:
        _add_section_heading(document, "Professional summary", spec)
        paragraph = document.add_paragraph(style="Resume Body")
        _font(paragraph.add_run(summary), spec)

    sections = draft.get("sections")
    if not isinstance(sections, dict):
        sections = {}
    for key, title_value in (
        ("experience", "Experience"),
        ("projects", "Projects"),
        ("education", "Education"),
        ("skills", "Skills"),
        ("languages", "Languages"),
        ("additional", "Additional information"),
    ):
        lines = _clean_lines(sections.get(key))
        if lines:
            _add_section_heading(document, title_value, spec)
            _add_lines(document, lines, spec)

    custom_sections = draft.get("custom_sections")
    if isinstance(custom_sections, list):
        for custom in custom_sections:
            if not isinstance(custom, dict):
                continue
            custom_title = str(custom.get("title") or "").strip()
            lines = _clean_lines(custom.get("lines"))
            if custom_title and lines:
                _add_section_heading(document, custom_title, spec)
                _add_lines(document, lines, spec)

    properties = document.core_properties
    properties.title = title
    properties.subject = "ATS-friendly resume"
    properties.comments = description
    properties.keywords = "resume, ATS, Axelyn Forge"
    document.save(output)
    return spec.id, TEMPLATE_VERSION
